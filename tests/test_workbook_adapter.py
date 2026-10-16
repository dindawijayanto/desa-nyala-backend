from pathlib import Path
from openpyxl import load_workbook

from app.schemas.input import OptimizationRequest
from app.services.offgridders_adapter import _configure_workbook
from tests.test_validation import sample_payload


def _find_value(ws, key, key_col=1, start_row=1):
    for r in range(start_row, ws.max_row + 1):
        if str(ws.cell(r, key_col).value or '').strip() == key:
            return r
    raise AssertionError(f"Missing key {key}")


def test_workbook_adapter_maps_three_scenarios_and_disables_wind(tmp_path):
    backend_root = Path(__file__).resolve().parents[1]
    template = backend_root / 'templates' / 'offgridders_input_template.xlsx'
    request = OptimizationRequest.model_validate(sample_payload())
    workbook_path = tmp_path / 'config' / 'input.xlsx'
    _configure_workbook(template, workbook_path, tmp_path, request)

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    settings = wb['settings']
    row = _find_value(settings, 'evaluated_days', key_col=2, start_row=11)
    assert settings.cell(row, 3).value == 365
    row = _find_value(settings, 'save_to_csv_flows_electricity_mg', key_col=2, start_row=11)
    assert settings.cell(row, 3).value == 'True'

    cases = wb['case_definitions']
    case_cols = {cases.cell(17, c).value: c for c in range(2, cases.max_column + 1) if cases.cell(17, c).value}
    perform_row = _find_value(cases, 'perform_simulation', key_col=1, start_row=18)
    assert cases.cell(perform_row, case_cols['diesel_mg']).value == 'True'
    assert cases.cell(perform_row, case_cols['pv-diesel-mg']).value == 'True'
    assert cases.cell(perform_row, case_cols['pv-diesel-storage-mg']).value == 'True'
    assert cases.cell(perform_row, case_cols['pv-wind-storage-diesel-mg']).value == 'False'
    wind_row = _find_value(cases, 'capacity_wind_kW', key_col=1, start_row=18)
    for case_name in ('diesel_mg', 'pv-diesel-mg', 'pv-diesel-storage-mg', 'pv-wind-storage-diesel-mg'):
        assert cases.cell(wind_row, case_cols[case_name]).value == 'None'
    wb.close()


def test_csv_input_is_not_required_and_generated_profile_files_are_local(tmp_path):
    from app.services.offgridders_adapter import _write_profiles
    from app.services.demand_profile import build_hourly_demand_profile
    request = OptimizationRequest.model_validate(sample_payload())
    path = _write_profiles(tmp_path, request)
    assert path.is_file()
    lines = path.read_text(encoding='utf-8').splitlines()
    assert lines[0] == 'Demand;SolarGen;Wind'
    assert len(lines) == request.project.simulation_days * 24 + 1
    assert all(line.endswith(';0.00000000') for line in lines[1:])
    daily = build_hourly_demand_profile(request.demand, 1)
    assert len(daily) == 24
