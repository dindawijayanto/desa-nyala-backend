#!/usr/bin/env python3
"""Release gate for DesaNyala's latest Offgridders API smoke-test run.

Run from the backend directory after updating runs/api_smoke_response.json:
    .venv\\Scripts\\python.exe validate_desanyala_release.py

This is an integration audit, not proof that a physical microgrid design is field-validated.
"""
from __future__ import annotations

import csv
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

SCENARIO_CASES = {
    "diesel_only": "diesel_mg",
    "pv_diesel": "pv-diesel-mg",
    "pv_diesel_battery": "pv-diesel-storage-mg",
}
STABILITY_PHRASE = "Stability criterion is strictly not fullfilled"
SHORTAGE_COMMENT = "Shortage penalty costs used in OEM not included in LCOE"
TOL = 1e-4


def fail(message: str) -> None:
    raise SystemExit(f"FAIL: {message}")


def number(row: dict[str, Any], key: str) -> float:
    try:
        value = float(row[key])
    except (KeyError, TypeError, ValueError):
        fail(f"Kolom {key!r} tidak tersedia atau bukan angka pada results.csv")
    if not math.isfinite(value):
        fail(f"Nilai {key!r} bukan angka finite: {value}")
    return value


def truthy_setting(workbook_path: Path, sheet_name: str, key: str, key_col: int, value_col: int, start_row: int) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        fail("openpyxl tidak tersedia pada .venv")
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for row_no in range(start_row, ws.max_row + 1):
            value = ws.cell(row_no, key_col).value
            if isinstance(value, str) and value.strip().casefold() == key.casefold():
                return str(ws.cell(row_no, value_col).value).strip().casefold()
    finally:
        wb.close()
    fail(f"Parameter {key!r} tidak ditemukan di sheet {sheet_name!r}")


def main() -> None:
    response_path = Path("runs/api_smoke_response.json")
    payload_path = Path("runs/api_smoke_payload.json")
    if not response_path.is_file() or not payload_path.is_file():
        fail("Jalankan smoke test API dahulu sehingga runs/api_smoke_response.json ada")

    try:
        response = json.loads(response_path.read_text(encoding="utf-8"))
        payload = json.loads(payload_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"Gagal membaca payload/response JSON: {exc}")

    scenarios = response.get("scenario_results")
    if not isinstance(scenarios, list) or not scenarios:
        fail("scenario_results kosong atau bukan list")
    by_id = {s.get("scenario_id"): s for s in scenarios if isinstance(s, dict)}
    requested = payload.get("optimization", {}).get("scenarios", list(SCENARIO_CASES))
    if set(by_id) != set(requested):
        fail(f"Skenario response {sorted(by_id)} tidak cocok dengan request {sorted(requested)}")
    if response.get("status") not in {"optimal", "feasible"}:
        fail(f"Status keseluruhan tidak dapat diterima untuk smoke test: {response.get('status')!r}")

    global_validation = response.get("validation") or {}
    for key in ("energy_balance_passed", "reliability_passed", "capacity_constraints_passed"):
        if global_validation.get(key) is not True:
            fail(f"Validasi global {key} bukan True")

    warning_text = "\n".join(str(x) for x in response.get("warnings", []))
    if "diulang" not in warning_text.casefold() or "sintetis" not in warning_text.casefold():
        fail("Respons tidak menjelaskan keterbatasan profil demand/PV sintetis")
    if not response.get("generated_at"):
        fail("generated_at tidak ada pada response")

    solver_message = str(global_validation.get("solver_message", ""))
    match = re.search(r"Data run:\s*(.+?)\s*$", solver_message, re.MULTILINE)
    if not match:
        fail("validation.solver_message tidak menyertakan lokasi folder output run")
    output_dir = Path(match.group(1).strip().strip('"'))
    if not output_dir.is_dir():
        fail(f"Folder output run tidak ditemukan: {output_dir}")
    run_dir = output_dir.parent
    results_path = output_dir / "results.csv"
    native_log = output_dir / "micro_grid_design_logfile.log"
    stdout_log = run_dir / "stdout.log"
    stderr_log = run_dir / "stderr.log"
    workbook_path = run_dir / "config" / "desanyala_input.xlsx"
    for path in (results_path, native_log, stdout_log, stderr_log, workbook_path):
        if not path.is_file():
            fail(f"Artefak run tidak ditemukan: {path}")

    with results_path.open("r", encoding="utf-8-sig", newline="") as handle:
        native_rows = list(csv.DictReader(handle))
    native_by_case = {str(row.get("case", "")).strip(): row for row in native_rows}
    for scenario_id in by_id:
        native_case = SCENARIO_CASES.get(scenario_id)
        if not native_case or native_case not in native_by_case:
            fail(f"Tidak ada baris results.csv untuk skenario {scenario_id!r}")

    # 1) Stability log must be scoped to the matching scenario only.
    log_text = native_log.read_text(encoding="utf-8", errors="replace")
    current_case: str | None = None
    unstable_cases: set[str] = set()
    for line in log_text.splitlines():
        started = re.search(r"Starting simulation of case\s+([A-Za-z0-9_-]+)", line, re.IGNORECASE)
        if started:
            current_case = started.group(1)
        if STABILITY_PHRASE in line and current_case:
            unstable_cases.add(current_case)

    # 2) Independent numerical checks plus backend validation flags.
    reliability = payload.get("reliability", {})
    shortage_policy = str(reliability.get("shortage_policy", "")).casefold()
    pv = payload.get("pv", {})
    battery = payload.get("battery", {})
    generators = payload.get("generators", {})
    for scenario_id, result in by_id.items():
        native_case = SCENARIO_CASES[scenario_id]
        row = native_by_case[native_case]
        validation = result.get("validation") or {}
        for key in ("energy_balance_passed", "reliability_passed", "capacity_constraints_passed"):
            if validation.get(key) is not True:
                fail(f"{scenario_id}: validation.{key} bukan True")

        lcoe = number(row, "lcoe")
        if lcoe <= 0:
            fail(f"{scenario_id}: LCOE harus positif, ditemukan {lcoe}")
        demand = number(row, "total_demand_annual_kWh")
        supplied = number(row, "total_demand_supplied_annual_kWh")
        shortage = number(row, "total_demand_shortage_annual_kWh")
        if shortage < -TOL or supplied < -TOL or supplied > demand + TOL:
            fail(f"{scenario_id}: nilai demand/supply/shortage tidak konsisten")
        if shortage_policy == "forbid" and (abs(shortage) > TOL or abs(supplied - demand) > TOL):
            fail(f"{scenario_id}: policy forbid tetapi shortage={shortage}, demand={demand}, supplied={supplied}")
        shortage_cost = number(row, "expenditures_shortage_annual")
        if abs(shortage_cost) > TOL and abs(shortage) <= TOL:
            fail(f"{scenario_id}: biaya shortage nonzero padahal shortage nol")

        if pv.get("enabled") and scenario_id in {"pv_diesel", "pv_diesel_battery"}:
            pv_cap = number(row, "capacity_pv_kWp")
            pv_min = float(pv.get("capacity_min_kwp", 0) or 0)
            pv_max = float(pv.get("capacity_max_kwp", math.inf) or math.inf)
            if pv_cap < pv_min - TOL or pv_cap > pv_max + TOL:
                fail(f"{scenario_id}: PV {pv_cap} kWp di luar batas [{pv_min}, {pv_max}]")
            if str(pv.get("capacity_mode", "")).casefold() == "fixed" and pv.get("fixed_capacity_kwp") is not None:
                if abs(pv_cap - float(pv["fixed_capacity_kwp"])) > 1e-3:
                    fail(f"{scenario_id}: kapasitas PV tidak sama dengan fixed_capacity_kwp")

        if battery.get("enabled") and scenario_id == "pv_diesel_battery":
            batt_cap = number(row, "capacity_storage_kWh")
            batt_power = number(row, "power_storage_kW")
            batt_min = float(battery.get("capacity_min_kwh", 0) or 0)
            batt_max = float(battery.get("capacity_max_kwh", math.inf) or math.inf)
            batt_power_max = float(battery.get("power_max_kw", math.inf) or math.inf)
            if batt_cap < batt_min - TOL or batt_cap > batt_max + TOL:
                fail(f"{scenario_id}: kapasitas baterai {batt_cap} kWh di luar [{batt_min}, {batt_max}]")
            if batt_power > batt_power_max + TOL:
                fail(f"{scenario_id}: daya baterai {batt_power} kW melebihi batas {batt_power_max} kW")

        if float(reliability.get("reserve_margin_fraction", 0) or 0) > 0 and generators.get("mode") == "new":
            genset_cap = number(row, "capacity_genset_kW")
            peak = number(row, "demand_peak_kW")
            target = peak * (1 + float(reliability["reserve_margin_fraction"]))
            if genset_cap + TOL < target:
                fail(f"{scenario_id}: genset {genset_cap} kW di bawah target reserve margin {target} kW")

        scenario_warnings = "\n".join(str(x) for x in validation.get("warnings", []))
        has_stability_warning = "stabilitas" in scenario_warnings.casefold() and "tidak" in scenario_warnings.casefold()
        if native_case in unstable_cases:
            if result.get("status") == "optimal" or not has_stability_warning:
                fail(f"{scenario_id}: native log mencatat deviasi stabilitas tetapi status/warning API tidak mencerminkannya")
        elif has_stability_warning:
            fail(f"{scenario_id}: warning stabilitas dari skenario lain bocor ke skenario ini")

    # 3) Economics config: include shortage penalties in LCOE and don't silently restore template defaults.
    flag_value = truthy_setting(workbook_path, "settings", "include_shortage_penalty_costs_in_lcoe", 2, 3, 11)
    if flag_value not in {"true", "1"}:
        fail(f"include_shortage_penalty_costs_in_lcoe bukan True pada workbook aktual: {flag_value!r}")
    shortage_penalty = truthy_setting(workbook_path, "input_constant", "shortage_penalty_costs", 1, 2, 6)
    try:
        penalty_value = float(shortage_penalty)
    except ValueError:
        fail(f"shortage_penalty_costs bukan angka: {shortage_penalty!r}")
    if not math.isfinite(penalty_value) or penalty_value < 0:
        fail(f"shortage_penalty_costs invalid: {penalty_value}")
    all_comments = " ".join(str(row.get("comments", "")) for row in native_rows)
    if SHORTAGE_COMMENT in all_comments:
        fail("CSV masih menyebut shortage penalty tidak dimasukkan ke LCOE")

    print("PASS: integration release gate")
    print(f"API status: {response.get('status')}; selected: {response.get('selected_scenario_id')}")
    print(f"Scenarios checked: {', '.join(by_id)}")
    print(f"Unstable native cases: {', '.join(sorted(unstable_cases)) if unstable_cases else 'none'}")
    print("Checked: per-scenario stability scope, validation flags, capacities, reserve margin, shortage, LCOE, workbook economics settings, profile caveats, and run logs.")
    print(f"Run trace: {run_dir}")
    print("NOTE: PASS means backend integration checks passed; it does not establish field readiness or validate synthetic profiles against real site data.")


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as exc:
        print(f"FAIL: unexpected validation error: {type(exc).__name__}: {exc}", file=sys.stderr)
        raise SystemExit(2)
